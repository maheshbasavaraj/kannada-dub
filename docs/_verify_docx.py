"""Verify the generated docx actually contains the expected content."""
from docx import Document

DOC = "docs/IMPLEMENTATION.docx"
doc = Document(DOC)

# collect body + table text
parts = [p.text for p in doc.paragraphs]
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            parts.append(c.text)
blob = "\n".join(parts)

checks = [
    "PYTHONPATH= TORCHDYNAMO_DISABLE",   # code block survived
    "291.71",                            # verified duration
    "364/364",                           # tensor load fix
    "2.18",                              # overflow figure
    "REF_SECONDS",                       # reference param
    "IndicF5", "rubberband", "demucs", "Whisper", "Ollama", "Vocos",
    "s4_voice", "narration.wav", "c__DMV_Clog_kn.mp4",
]
print(f"paragraphs   : {len(doc.paragraphs)}")
print(f"tables       : {len(doc.tables)}")
print(f"body chars   : {len(blob)}")
print("-" * 46)
for c in checks:
    n = blob.count(c)
    print(f"{'OK ' if n else 'MISS'} {c:<34} x{n}")

missing = [c for c in checks if c not in blob]
print("-" * 46)
print("RESULT:", "ALL PRESENT" if not missing else f"MISSING {missing}")

# code block newlines: the Code-styled paragraph should hold multiple lines
for p in doc.paragraphs:
    if p.style.name == "Code":
        print("code paragraph lines:", len(p.text.splitlines()))
